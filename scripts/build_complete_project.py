from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path('/Users/macbook/Downloads/PROJECT MAIN WORK CORRECTED.docx')
OUT = ROOT / 'deliverables'
OUT.mkdir(exist_ok=True)
DOCX = OUT / 'AMEH_JOSEPH_JUNIOR_COMPLETE_PROJECT.docx'


def set_cell_text(cell, text, bold=False):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = 'Times New Roman'
    r.font.size = Pt(10)


def shade(cell, fill='D9EAF7'):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.autofit = True
    for i, h in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], h, True)
        shade(table.rows[0].cells[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    return table


def field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    text = OxmlElement('w:instrText'); text.set(qn('xml:space'), 'preserve'); text.text = instruction
    separate = OxmlElement('w:fldChar'); separate.set(qn('w:fldCharType'), 'separate')
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, text, separate, end])


def set_page_number(section, roman=False):
    sect_pr = section._sectPr
    pg = sect_pr.find(qn('w:pgNumType'))
    if pg is None:
        pg = OxmlElement('w:pgNumType'); sect_pr.append(pg)
    pg.set(qn('w:start'), '1')
    pg.set(qn('w:fmt'), 'lowerRoman' if roman else 'decimal')
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field(p, 'PAGE')


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.add_run(text)
    return p


def body(doc, text, bold_lead=None):
    p = doc.add_paragraph(style='Body Text')
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.add_run(text)
    return p


def caption(doc, text):
    p = doc.add_paragraph(style='Caption')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)


def picture(doc, path, caption_text, width=6.0):
    if Path(path).exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        caption(doc, caption_text)


def title_page(doc):
    for _ in range(3): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('DEVELOPMENT OF A CONVOLUTIONAL NEURAL NETWORK MODEL FOR CATFISH DEVELOPMENTAL-STAGE CLASSIFICATION AND BIOMETRIC ESTIMATION')
    r.bold = True; r.font.size = Pt(16)
    for _ in range(2): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('BY\n\nAMEH JOSEPH JUNIOR\nPG/MSC/19/89745'); r.bold = True; r.font.size = Pt(14)
    for _ in range(3): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('A RESEARCH PROJECT SUBMITTED TO THE DEPARTMENT OF COMPUTER SCIENCE, UNIVERSITY OF NIGERIA, NSUKKA, IN PARTIAL FULFILMENT OF THE REQUIREMENTS FOR THE AWARD OF THE MASTER OF SCIENCE (M.Sc.) DEGREE IN COMPUTER SCIENCE').bold = True
    for _ in range(3): doc.add_paragraph()
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('SUPERVISOR: PROF. COLLINS UDANOR\n\nAUGUST 2026').bold = True
    doc.add_page_break()


def preliminaries(doc):
    heading(doc, 'DECLARATION', 1)
    body(doc, 'I, Ameh Joseph Junior (PG/MSC/19/89745), declare that this project is an account of the work documented in the accompanying repository and experimental notebook. Sources used in developing the study are acknowledged in the text and reference list. No result has been presented as independently verified where the available project evidence does not support such a claim.')
    body(doc, '\nSignature: __________________________    Date: __________________')
    doc.add_page_break()
    heading(doc, 'CERTIFICATION', 1)
    body(doc, 'This project has been presented for examination with the approval of the supervisor. The signature fields below are reserved for institutional completion.')
    body(doc, '\nProf. Collins Udanor (Supervisor)\nSignature: __________________________    Date: __________________')
    body(doc, '\nHead of Department\nSignature: __________________________    Date: __________________')
    doc.add_page_break()
    heading(doc, 'DEDICATION', 1)
    body(doc, 'This work is dedicated to the advancement of practical artificial intelligence for sustainable aquaculture and food production.')
    doc.add_page_break()
    heading(doc, 'ACKNOWLEDGEMENTS', 1)
    body(doc, 'I acknowledge the guidance of my supervisor, Prof. Collins Udanor, and the support of the Department of Computer Science, University of Nigeria, Nsukka. I also acknowledge everyone who contributed to the acquisition, measurement and organization of the catfish data used in this study. Any remaining errors are mine.')
    doc.add_page_break()
    heading(doc, 'ABSTRACT', 1)
    body(doc, 'This study developed and evaluated a transfer-learning system for classifying African catfish images into fingerling, juvenile and adult developmental stages while estimating standard length (SL), total length (TL) and weight. The archived notebook contains 357 valid images representing 117 independently identified fish. To prevent information leakage from repeated photographs of the same fish, splitting was performed by fish identity, producing 249 training images from 81 fish, 51 validation images from 17 fish and 57 test images from 19 fish. EfficientNetB0 was selected as the best experimental backbone. On the held-out test set, image-level accuracy was 82.46% and macro F1-score was 81.23%. Aggregating predictions by fish improved accuracy to 94.74% and macro F1-score to 94.34%. Biometric estimation was substantially weaker: fish-level mean absolute errors were 11.24 cm for SL, 15.64 cm for TL and 285.03 g for weight; the regression model did not outperform the class-mean baseline. A local Streamlit application was also implemented and tested with 31 automated tests. Its archived screenshots demonstrate image upload, validation, stage prediction and biometric display, but the saved application checkpoint is InceptionV3 and is therefore reported separately from the evaluated EfficientNetB0 experiment. The findings support fish-grouped stage classification as a promising proof of concept, while showing that reliable image-only biometric estimation requires more independent fish, calibrated image acquisition and a visible scale reference.')
    doc.add_page_break()
    heading(doc, 'TABLE OF CONTENTS', 1)
    p = doc.add_paragraph(); field(p, 'TOC \\o "1-3" \\h \\z \\u')
    for item in ['Chapter One: Introduction','Chapter Two: Literature Review','Chapter Three: System Analysis and Design','Chapter Four: System Implementation, Results and Discussion','Chapter Five: Summary, Conclusion and Recommendations','References','Appendices A–D']:
        body(doc, item)
    body(doc, 'Note: update the automatic field in Microsoft Word or LibreOffice to display final page numbers.')
    doc.add_page_break()
    heading(doc, 'LIST OF TABLES', 1)
    p = doc.add_paragraph(); field(p, 'TOC \\h \\z \\c "Table"')
    for item in ['Tables 4.1–4.10: Experimental configuration, dataset, split, metrics, regression and software tests','Table 5.1: Findings mapped to the study objectives','Appendix A table: Reproducibility and evidence checklist']:
        body(doc, item)
    doc.add_page_break()
    heading(doc, 'LIST OF FIGURES', 1)
    p = doc.add_paragraph(); field(p, 'TOC \\h \\z \\c "Figure"')
    for item in ['Figures 4.1–4.4: Dataset, confusion matrix, ROC and regression diagnostics','Figures 4.5–4.8: Local application interface and representative predictions']:
        body(doc, item)
    doc.add_page_break()
    heading(doc, 'LIST OF ABBREVIATIONS', 1)
    add_table(doc, ['Abbreviation', 'Meaning'], [
        ('AI','Artificial Intelligence'), ('CNN','Convolutional Neural Network'),
        ('CI','Confidence Interval'), ('F1','Harmonic mean of precision and recall'),
        ('Grad-CAM','Gradient-weighted Class Activation Mapping'), ('MAE','Mean Absolute Error'),
        ('RMSE','Root Mean Squared Error'), ('SL','Standard Length'), ('TL','Total Length'),
        ('UI','User Interface'), ('ViT','Vision Transformer')])
    doc.add_page_break()


def corrected_source_text(text, in_ch3):
    replacements = {
        'The aim of the study is to develop a highly accurate and efficient CNN-based classification model to classify catfish images into various categories representing different developmental stages.':
        'The aim of the study is to develop and evaluate a CNN-based system for classifying African catfish images into fingerling, juvenile and adult developmental stages and for estimating standard length, total length and weight from surface images.',
        'Deploy the developed model as a web-based application for practical use in catfish image classification and biometric estimation.':
        'Implement and demonstrate the developed prediction workflow in a locally executable web application for catfish image classification and biometric estimation.',
    }
    text = replacements.get(text, text)
    if in_ch3:
        text = text.replace('299 × 299', '224 × 224').replace('299×299', '224×224')
        text = text.replace('two continuous biometric parameters', 'three continuous biometric parameters')
        if text == 'Fish Length': text = 'Standard Length (SL) and Total Length (TL)'
        if text == 'Estimated body length': text = 'Estimated standard length and total length'
        text = text.replace('Mean Squared Error (MSE)', 'Huber loss')
        text = text.replace('separate frontend developed using HTML, CSS, and JavaScript', 'Streamlit user interface')
    return text


def import_chapters_1_to_3(doc):
    src = Document(SOURCE)
    in_ch3 = False
    for p in src.paragraphs:
        t = p.text.strip()
        if not t: continue
        if t == 'CHAPTER THREE': in_ch3 = True
        t = corrected_source_text(t, in_ch3)
        if t in {'CHAPTER ONE','CHAPTER TWO','CHAPTER THREE'}:
            if t != 'CHAPTER ONE': doc.add_page_break()
            heading(doc, t, 1); continue
        if t in {'INTRODUCTION','LITERATURE REVIEW','SYSTEM ANALYSIS AND DESIGN'}:
            h = heading(doc, t, 1); h.alignment = WD_ALIGN_PARAGRAPH.CENTER; continue
        if re.match(r'^\d+\.\d+(?:\.\d+){0,2}\s+', t):
            level = min(3, t.split()[0].count('.'))
            heading(doc, t, level); continue
        if len(t) < 75 and (t.endswith(':') or t in {'Aim of the Study','Objectives of the Study','Scope of the Study','Significance of the Study','AlexNet','Visual Geometry Group (VGG) Network','Residual Network (ResNet)','Densely Connected Convolutional Networks (DenseNet)','EfficientNet','Vision Transformers (ViTs)'}):
            heading(doc, t.rstrip(':'), 2); continue
        if len(t) < 80 and t in {'Fingerling','Juvenile','Adult','Fish Weight','Training dataset','Validation dataset','Testing dataset','Classification Accuracy','Precision','Recall','F1-Score','Confusion Matrix','Mean Absolute Error (MAE)','Root Mean Squared Error (RMSE)'}:
            bullet(doc, t); continue
        body(doc, t)


def chapter_four(doc):
    doc.add_page_break(); heading(doc, 'CHAPTER FOUR', 1); heading(doc, 'SYSTEM IMPLEMENTATION, RESULTS AND DISCUSSION', 1)
    heading(doc, '4.0 Introduction', 2)
    body(doc, 'This chapter reports only results preserved in the project notebook, repository and archived application screenshots. It distinguishes the EfficientNetB0 experimental model evaluated in the notebook from the InceptionV3 checkpoint used by the local application. This distinction prevents application demonstrations from being interpreted as held-out experimental evidence.')
    heading(doc, '4.1 System and Experimental Architecture', 2)
    body(doc, 'The experimental pipeline comprised dataset auditing, fish-identity grouping, stratified group splitting, augmentation, transfer learning, validation-based model selection, held-out image-level evaluation and fish-level aggregation. The multitask network shared an EfficientNetB0 feature extractor and used separate classification and regression heads. The local Streamlit application formed a separate inference interface around the available InceptionV3 checkpoint.')
    add_table(doc, ['Evidence stream','Backbone','Purpose','Status'], [
        ('Notebook experiment','EfficientNetB0','Research evaluation on grouped validation/test data','Metrics preserved; final .keras checkpoint absent'),
        ('Local application','InceptionV3','Interactive upload and demonstration','Weights and screenshots preserved'),
    ])
    caption(doc, 'Table 4.1: Separation of the experimental and application evidence streams')
    heading(doc, '4.2 Development Tools', 2)
    add_table(doc, ['Component','Technology','Role'], [
        ('Programming','Python','Data processing, model training and application logic'),
        ('Deep learning','TensorFlow/Keras','Transfer learning and multitask neural networks'),
        ('Data analysis','pandas, NumPy, scikit-learn','Metadata, grouping, scaling and metrics'),
        ('Image processing','Pillow, OpenCV','Image loading, validation and cropping experiments'),
        ('Visualization','Matplotlib, Seaborn','Charts and diagnostic plots'),
        ('Application','Streamlit','Local web user interface'),
        ('Testing','pytest','Automated software tests')])
    caption(doc, 'Table 4.2: Software tools used in the project')
    heading(doc, '4.3 Dataset Description and Audit', 2)
    body(doc, 'The archived audit reported 357 valid images representing 117 independent fish. Multiple photographs existed for some fish; consequently, the image is not the independent biological sampling unit. Two exact duplicate groups were found, but no duplicate crossed fish identities, classes or final data splits. Five empty folders and one unsupported file were also reported.')
    add_table(doc, ['Class','Images','Independent fish','Images per fish'], [
        ('Adult','104','33','3.15'), ('Fingerling','106','35','3.03'), ('Juvenile','147','49','3.00'), ('Total','357','117','3.05')])
    caption(doc, 'Table 4.3: Dataset composition by developmental stage')
    picture(doc, ROOT/'chapter_4_5_materials/images/cell_12_image_01.png', 'Figure 4.1: Dataset counts and exploratory distributions preserved in the notebook', 6.1)
    heading(doc, '4.4 Preprocessing and Leakage-Controlled Split', 2)
    body(doc, 'Images were resized to 224 × 224 pixels. Training-only augmentation included geometric and photometric transformations. Regression targets were scaled using statistics fitted on the training set: separate standard scalers for SL and TL, and log1p transformation followed by standard scaling for weight. The deterministic seed was 42.')
    add_table(doc, ['Subset','Images','Independent fish','Adult images/fish','Fingerling images/fish','Juvenile images/fish'], [
        ('Training','249','81','74 / 23','73 / 24','102 / 34'),
        ('Validation','51','17','15 / 5','15 / 5','21 / 7'),
        ('Test','57','19','15 / 5','18 / 6','24 / 8')])
    caption(doc, 'Table 4.4: Fish-grouped data split')
    body(doc, 'Leakage checks passed: no fish identity and no exact image hash occurred across training, validation and test partitions. This control was essential because randomly splitting photographs would allow different views of the same fish to occur in multiple subsets and inflate performance.')
    heading(doc, '4.5 Model Implementation and Training', 2)
    body(doc, 'EfficientNetB0, MobileNetV3Large and DenseNet121 were considered during backbone selection. EfficientNetB0 produced the best validation result. The multitask model combined a shared pretrained backbone with a developmental-stage softmax head and three regression outputs for SL, TL and weight. Classification used categorical cross-entropy; regression used Huber losses with task weights of 0.25 for SL, 0.25 for TL and 0.50 for weight, alongside a classification weight of 1.00. Adam was used for head training and AdamW for fine-tuning.')
    body(doc, 'The saved notebook run was explicitly a quick-mode experiment: two head-training epochs and one fine-tuning epoch, batch size 16. Cross-validation was disabled. Therefore, the reported metrics describe the archived run and should not be interpreted as a fully converged or cross-validated estimate.')
    add_table(doc, ['Validation metric','Value'], [('Accuracy','92.16%'),('Balanced accuracy','92.38%'),('Macro F1-score','92.38%'),('Fingerling recall','93.33%'),('Juvenile recall','90.48%'),('Adult recall','93.33%')])
    caption(doc, 'Table 4.5: Best validation classification results')
    heading(doc, '4.6 Held-Out Classification Results', 2)
    body(doc, 'On the 57-image test set, accuracy was 82.46% and macro F1-score was 81.23%. When the predictions from repeated views were aggregated at fish level, accuracy increased to 94.74% and macro F1-score to 94.34%. Fish-level results are biologically more appropriate because they evaluate the 19 independent fish rather than treating 57 correlated views as independent observations. Nevertheless, the fish-level estimate is based on a small test sample and has a wide uncertainty interval.')
    add_table(doc, ['Metric','Image level (57 images)','Fish level (19 fish)'], [
        ('Accuracy','82.46%','94.74%'),('Balanced accuracy','80.19%','93.33%'),
        ('Macro precision','83.93%','96.30%'),('Macro recall','80.19%','93.33%'),
        ('Macro F1-score','81.23%','94.34%'),('Weighted F1-score','82.05%','94.60%'),
        ('Juvenile recall','91.67%','100.00%')])
    caption(doc, 'Table 4.6: Held-out classification performance')
    add_table(doc, ['Fish-level statistic','Bootstrap 95% confidence interval'], [('Accuracy','84.21%–100.00%'),('Macro F1-score','75.93%–100.00%'),('Juvenile recall','100.00%–100.00%')])
    caption(doc, 'Table 4.7: Bootstrap uncertainty for fish-level classification')
    picture(doc, ROOT/'chapter_4_5_materials/images/cell_28_image_08.png', 'Figure 4.2: Fish-level confusion matrices', 6.0)
    picture(doc, ROOT/'chapter_4_5_materials/images/cell_28_image_09.png', 'Figure 4.3: Receiver-operating-characteristic curves', 6.0)
    heading(doc, '4.7 Biometric Estimation Results', 2)
    body(doc, 'Biometric estimation was not sufficiently accurate for dependable field use. The fish-level model did not outperform the class-mean baseline. Weight prediction was particularly unstable, with negative R² and very high percentage error. These outcomes are consistent with the absence of a visible scale reference and with changes in camera distance and background, which prevent absolute physical size from being uniquely inferred from pixels.')
    add_table(doc, ['Target','MAE','RMSE','R²','MAPE','Bootstrap 95% CI for MAE'], [
        ('SL','11.24 cm','14.10 cm','0.286','64.49%','7.50–15.25 cm'),
        ('TL','15.64 cm','17.45 cm','0.163','94.40%','12.27–19.27 cm'),
        ('Weight','285.03 g','519.33 g','−0.156','281.30%','87.50–491.18 g')])
    caption(doc, 'Table 4.8: Fish-level biometric estimation performance')
    add_table(doc, ['Class','SL MAE','TL MAE','Weight MAE'], [('Adult','25.08 cm','27.93 cm','1003.44 g'),('Fingerling','7.19 cm','13.45 cm','33.46 g'),('Juvenile','5.62 cm','9.60 cm','24.71 g')])
    caption(doc, 'Table 4.9: Class-specific biometric errors')
    picture(doc, ROOT/'chapter_4_5_materials/images/cell_28_image_11.png', 'Figure 4.4: Regression diagnostic plots preserved in the notebook', 6.1)
    heading(doc, '4.8 Explainability Attempt', 2)
    body(doc, 'The notebook attempted to generate six Grad-CAM examples, but all attempts failed because the gradcam function was called with an unsupported class_index keyword argument. The archived output records 0 successful images out of 6. Consequently, no Grad-CAM visual is presented as a successful explainability result. Repairing and rerunning this analysis is future work.')
    heading(doc, '4.9 Local Application Implementation', 2)
    body(doc, 'The Streamlit application supports image upload and webcam input, accepts JPEG and PNG files up to 10 MB, checks for unreadable, unusually small, blank, dark, bright or blurred inputs, performs a generic fish-screening check and displays the predicted stage, confidence and biometric estimates. It also reports plausibility warnings. The application has no verified public deployment, database, authentication or prediction-history subsystem; it is therefore described as a local implementation.')
    picture(doc, ROOT/'chapter_4_5_materials/app_screenshots/02_app_full_page.png', 'Figure 4.5: Full local Streamlit application interface', 5.5)
    heading(doc, '4.10 Representative Application Predictions', 2)
    body(doc, 'The archived screenshots show outputs for the three developmental-stage labels. They are interface demonstrations generated by the InceptionV3 application checkpoint and are not additional observations from the EfficientNetB0 held-out test set. The adult example produced SL greater than TL, and the application correctly displayed a plausibility warning; this illustrates the practical importance of output validation.')
    picture(doc, ROOT/'chapter_4_5_materials/app_screenshots/03_fingerling_prediction_result.png', 'Figure 4.6: Fingerling application prediction example', 5.4)
    picture(doc, ROOT/'chapter_4_5_materials/app_screenshots/04_juvenile_prediction_result.png', 'Figure 4.7: Juvenile application prediction example', 5.4)
    picture(doc, ROOT/'chapter_4_5_materials/app_screenshots/05_adult_prediction_result.png', 'Figure 4.8: Adult application prediction with biometric plausibility warning', 5.4)
    heading(doc, '4.11 Software Testing', 2)
    add_table(doc, ['Test area','Evidence','Outcome'], [
        ('Image validation','Automated tests for file and image-quality checks','Passed'),
        ('Preprocessing','Automated tests for input preparation','Passed'),
        ('Prediction logic','Automated tests for output handling and plausibility','Passed'),
        ('Analysis utilities','Automated tests for analysis functions','Passed'),
        ('Complete automated suite','31 pytest tests','31 passed')])
    caption(doc, 'Table 4.10: Software test summary')
    heading(doc, '4.12 Discussion of Findings', 2)
    body(doc, 'The results demonstrate that transfer learning can distinguish developmental stages in the collected data, particularly when several views are combined for each fish. The increase from 82.46% image-level accuracy to 94.74% fish-level accuracy suggests that aggregating repeated observations can reduce view-specific uncertainty. However, only 19 independent fish were available in the test set, so the apparent improvement requires confirmation on a larger external cohort.')
    body(doc, 'The regression results do not support claims of accurate image-only length or weight measurement. Classification can exploit shape, texture and contextual differences associated with development, whereas absolute physical measurements require scale. Camera-to-fish distance, perspective and background can make different-sized fish occupy similar pixel dimensions. The adult errors and negative weight R² reinforce this limitation.')
    body(doc, 'Finally, the model mismatch affects reproducibility. The evaluated EfficientNetB0 checkpoint and preprocessing artifacts were not preserved in the repository, whereas an InceptionV3 weights file was preserved for the application. A future release should export the selected model, scalers, class map and a machine-readable test-prediction file as a single versioned bundle.')
    heading(doc, '4.13 Chapter Summary', 2)
    body(doc, 'This chapter reported leakage-controlled classification results, biometric errors, application demonstrations and software tests. The study provides promising stage-classification evidence, but the biometric component, Grad-CAM analysis and public deployment objective remain incomplete.')


def chapter_five(doc):
    doc.add_page_break(); heading(doc, 'CHAPTER FIVE', 1); heading(doc, 'SUMMARY, CONCLUSION AND RECOMMENDATIONS', 1)
    heading(doc, '5.0 Introduction', 2)
    body(doc, 'This chapter summarizes the study, relates the verified evidence to the objectives, presents the conclusion and states the contributions, limitations and recommendations.')
    heading(doc, '5.1 Summary of the Study', 2)
    body(doc, 'The study investigated whether transfer learning could classify African catfish images into fingerling, juvenile and adult stages and estimate SL, TL and weight. A fish-grouped evaluation design was used to prevent repeated photographs of the same fish from crossing data partitions. The preserved notebook evaluated an EfficientNetB0 multitask model, while a separate InceptionV3 checkpoint supported a local Streamlit demonstration application.')
    heading(doc, '5.2 Summary of Major Findings', 2)
    add_table(doc, ['Objective','Verified finding','Status'], [
        ('Collect and preprocess a catfish dataset','357 valid images from 117 fish were audited and prepared.','Achieved in archived run; original input archive is not in repository'),
        ('Train a stage classifier and biometric estimator','EfficientNetB0 was trained; fish-level macro F1 was 94.34%, but regression did not beat the class-mean baseline.','Partly achieved'),
        ('Develop a user-friendly web application','A local Streamlit app supports upload/webcam, validation and results display.','Achieved locally'),
        ('Evaluate classification and regression','Image- and fish-level metrics, confidence intervals and regression errors were recorded.','Achieved; no independent expert comparison study'),
        ('Deploy the application','No public URL or hosting evidence exists.','Not achieved; local implementation only')])
    caption(doc, 'Table 5.1: Findings mapped to the study objectives')
    heading(doc, '5.3 Conclusion', 2)
    body(doc, 'The study concludes that CNN transfer learning is a promising approach for classifying the developmental stage of African catfish from photographs when evaluation is performed at the independent-fish level. The archived experiment achieved 94.74% fish-level accuracy, although the small 19-fish test set limits precision and external validity. The evidence does not support dependable image-only estimates of absolute length or weight: all regression errors were high, weight R² was negative and the model failed to outperform the class-mean baseline. Therefore, the application should presently be treated as a research prototype for stage classification and interface evaluation, not as a validated biometric measurement or production deployment system.')
    heading(doc, '5.4 Contributions of the Study', 2)
    bullet(doc, 'A locally collected dataset structure that explicitly identifies 117 independent fish across three developmental stages.')
    bullet(doc, 'A leakage-controlled evaluation protocol that groups repeated photographs by fish identity.')
    bullet(doc, 'Direct comparison of image-level and fish-level classification, showing the value of multi-view aggregation.')
    bullet(doc, 'A multitask proof of concept that reports both successes and failure modes for classification and biometric estimation.')
    bullet(doc, 'A tested local Streamlit interface with image-quality checks and biometric plausibility warnings.')
    heading(doc, '5.5 Limitations', 2)
    bullet(doc, 'Only 117 independent fish were available, with 19 fish in the test set and wide bootstrap intervals.')
    bullet(doc, 'The saved run used quick-mode training for three total epochs and did not perform cross-validation.')
    bullet(doc, 'Photographs lacked a fixed acquisition geometry or visible scale reference, confounding absolute size with camera distance and perspective.')
    bullet(doc, 'The evaluated EfficientNetB0 checkpoint and scalers were not preserved, limiting exact reproduction of notebook predictions.')
    bullet(doc, 'The application uses a different InceptionV3 checkpoint; its screenshots are demonstrations rather than held-out EfficientNetB0 evidence.')
    bullet(doc, 'Grad-CAM generation failed, so model attention was not successfully verified.')
    bullet(doc, 'No external-site validation, independent expert-comparison protocol or public deployment evidence was available.')
    heading(doc, '5.6 Recommendations', 2)
    bullet(doc, 'Collect a substantially larger dataset from multiple farms, cameras, operators, backgrounds and seasons, retaining unique fish identifiers.')
    bullet(doc, 'Use a calibration board, ruler or known-size marker in every photograph and standardize camera distance and orientation.')
    bullet(doc, 'Evaluate on an untouched external farm dataset and report fish-level confidence intervals as the primary biological analysis.')
    bullet(doc, 'Train to convergence with early stopping and repeated grouped cross-validation, followed by a single locked test evaluation.')
    bullet(doc, 'Export the selected model, preprocessing scalers, label map, environment lock file and test predictions as one versioned artifact.')
    bullet(doc, 'Repair Grad-CAM and include sanity checks before making explainability claims.')
    bullet(doc, 'Unify the application and research model, then perform usability, latency and reliability testing before public deployment.')
    heading(doc, '5.7 Suggestions for Future Work', 2)
    body(doc, 'Future work should investigate segmentation-based length measurement, calibrated monocular geometry, stereo or depth cameras, and multi-view fusion. Ordinal-stage modelling may better reflect the progressive nature of growth than an unconstrained three-class output. For weight estimation, future studies should combine calibrated length, body-width and condition-factor features rather than depend on an unscaled photograph alone. A prospective field study should compare predictions against standardized manual measurements made by trained fisheries personnel.')


def references_and_appendices(doc):
    doc.add_page_break(); heading(doc, 'REFERENCES', 1)
    refs = [
        'Ahmad, U., Ali, M. J., Khan, F. A., Khan, A. A., Rehman, A. U., Shahid, M. M. A., Haq, M. A., Khan, I., & Zamil, S. (2023). Large scale fish images classification and localization using transfer learning and localization-aware CNN architecture. Computer Systems Science and Engineering, 45(2), 2125–2140. https://doi.org/10.32604/csse.2023.031008',
        'Deka, J., Laskar, S., & Baklial, B. (2023). Automated freshwater fish species classification using deep CNN. Journal of The Institution of Engineers (India): Series B, 104(3), 603–621. https://doi.org/10.1007/s40031-023-00883-2',
        'Dosovitskiy, A., Beyer, L., Kolesnikov, A., Weissenborn, D., Zhai, X., Unterthiner, T., et al. (2021). An image is worth 16×16 words: Transformers for image recognition at scale. International Conference on Learning Representations.',
        'Goodfellow, I., Bengio, Y., & Courville, A. (2016). Deep learning. MIT Press.',
        'Hasegawa, T., Kondo, K., & Senou, H. (2024). Transferable deep learning model for the identification of fish species for various fishing grounds. Journal of Marine Science and Engineering, 12(3).',
        'He, K., Zhang, X., Ren, S., & Sun, J. (2016). Deep residual learning for image recognition. Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition, 770–778.',
        'Huang, G., Liu, Z., van der Maaten, L., & Weinberger, K. Q. (2017). Densely connected convolutional networks. Proceedings of the IEEE Conference on Computer Vision and Pattern Recognition, 4700–4708.',
        'Krizhevsky, A., Sutskever, I., & Hinton, G. E. (2012). ImageNet classification with deep convolutional neural networks. Advances in Neural Information Processing Systems, 25.',
        'LeCun, Y., Bottou, L., Bengio, Y., & Haffner, P. (1998). Gradient-based learning applied to document recognition. Proceedings of the IEEE, 86(11), 2278–2324.',
        'Pan, S. J., & Yang, Q. (2010). A survey on transfer learning. IEEE Transactions on Knowledge and Data Engineering, 22(10), 1345–1359.',
        'Selvaraju, R. R., Cogswell, M., Das, A., Vedantam, R., Parikh, D., & Batra, D. (2017). Grad-CAM: Visual explanations from deep networks via gradient-based localization. Proceedings of the IEEE International Conference on Computer Vision, 618–626.',
        'Shorten, C., & Khoshgoftaar, T. M. (2019). A survey on image data augmentation for deep learning. Journal of Big Data, 6, 60.',
        'Simonyan, K., & Zisserman, A. (2015). Very deep convolutional networks for large-scale image recognition. International Conference on Learning Representations.',
        'Tan, M., & Le, Q. V. (2019). EfficientNet: Rethinking model scaling for convolutional neural networks. Proceedings of the 36th International Conference on Machine Learning, 6105–6114.',
        'Ameh, J. J. (2026). Catfish attribute prediction [Computer software and research repository]. GitHub. https://github.com/bechosen-spec/catfish-attribute-prediction',
    ]
    for ref in refs:
        p = doc.add_paragraph(style='Reference'); p.add_run(ref)
    body(doc, 'Editorial note: several recent empirical citations retained from the supplied Chapter Two manuscript require final verification against the author’s source copies before institutional submission; incomplete bibliographic records were not invented in this compilation.')
    doc.add_page_break(); heading(doc, 'APPENDIX A: REPRODUCIBILITY AND EVIDENCE CHECKLIST', 1)
    add_table(doc, ['Artifact','Available','Interpretation'], [
        ('Original notebook with embedded outputs','Yes','Supports archived metrics and figures'),
        ('Original dataset ZIP and spreadsheets','No','Dataset cannot be reconstructed from repository alone'),
        ('EfficientNetB0 final .keras checkpoint','No','Exact notebook inference cannot be rerun'),
        ('InceptionV3 application weights','Yes','Supports local application inference'),
        ('Regression scalers from notebook run','No','Required for exact biometric reproduction'),
        ('Application screenshots','Yes','Supports interface description only'),
        ('Automated tests','Yes','31 tests passed during audit'),
        ('Public deployment URL','No','No deployment claim made'),
        ('Repository URL','Yes','https://github.com/bechosen-spec/catfish-attribute-prediction')])
    doc.add_page_break(); heading(doc, 'APPENDIX B: LOCAL APPLICATION USER GUIDE', 1)
    body(doc, '1. Create a Python environment and install the packages listed in requirements.txt.\n2. Ensure InceptionV3_best_model.weights.h5 is located in the repository root.\n3. Start the interface with: streamlit run app.py\n4. Upload a clear JPEG or PNG catfish image (maximum 10 MB) or use the webcam option.\n5. Review the predicted stage, confidence, SL, TL and weight values. Treat biometric values as experimental estimates and observe any plausibility warning.\n6. Do not use the current prototype for commercial grading or biological measurement without external validation and image calibration.')
    doc.add_page_break(); heading(doc, 'APPENDIX C: REPOSITORY STRUCTURE', 1)
    body(doc, 'Key files include app.py (Streamlit entry point), src/model.py (application architecture), src/prediction.py (inference), src/image_validator.py (input checks), catfish_multitask_colab.ipynb (archived experiment), InceptionV3_best_model.weights.h5 (application checkpoint), tests/ (automated tests), and chapter_4_5_materials/ (extracted notebook outputs and application screenshots).')
    doc.add_page_break(); heading(doc, 'APPENDIX D: AUDIT TRAIL FOR REPORTED RESULTS', 1)
    body(doc, 'All numerical results in Chapters Four and Five were transcribed from the embedded outputs of catfish_multitask_colab.ipynb and the extracted materials in chapter_4_5_materials. The application description was checked against the repository source and archived screenshots. The original Chapters One to Three were retained as the manuscript foundation, with targeted corrections for the aim, local-implementation objective, 224 × 224 input resolution, SL/TL terminology, Huber regression loss and Streamlit implementation.')


def style_document(doc):
    styles = doc.styles
    normal = styles['Normal']; normal.font.name='Times New Roman'; normal.font.size=Pt(12)
    normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    for sname in ['Body Text','List Bullet']:
        s=styles[sname]; s.font.name='Times New Roman'; s.font.size=Pt(12)
        s.paragraph_format.line_spacing=1.5; s.paragraph_format.space_after=Pt(6)
        if sname=='Body Text': s.paragraph_format.first_line_indent=Inches(.5); s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for i in (1,2,3):
        s=styles[f'Heading {i}']; s.font.name='Times New Roman'; s.font.color.rgb=RGBColor(0,0,0); s.font.bold=True
        s.font.size=Pt(14 if i==1 else 12); s.paragraph_format.keep_with_next=True
        s.paragraph_format.space_before=Pt(12); s.paragraph_format.space_after=Pt(6)
        if i==1: s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
    cap=styles['Caption']; cap.font.name='Times New Roman'; cap.font.size=Pt(10); cap.font.italic=True; cap.font.color.rgb=RGBColor(0,0,0)
    if 'Reference' not in styles:
        rs=styles.add_style('Reference', WD_STYLE_TYPE.PARAGRAPH)
    rs=styles['Reference']; rs.font.name='Times New Roman'; rs.font.size=Pt(11); rs.paragraph_format.left_indent=Inches(.5); rs.paragraph_format.first_line_indent=Inches(-.5); rs.paragraph_format.space_after=Pt(6)
    for sec in doc.sections:
        sec.top_margin=Inches(1); sec.bottom_margin=Inches(1); sec.left_margin=Inches(1.25); sec.right_margin=Inches(1)


def main():
    doc=Document(); style_document(doc)
    title_page(doc); preliminaries(doc)
    front=doc.sections[0]; set_page_number(front, roman=True)
    mainsec=doc.add_section(WD_SECTION.NEW_PAGE); set_page_number(mainsec, roman=False)
    import_chapters_1_to_3(doc); chapter_four(doc); chapter_five(doc); references_and_appendices(doc)
    style_document(doc)
    props=doc.core_properties; props.title='Catfish Developmental-Stage Classification and Biometric Estimation'; props.author='Ameh Joseph Junior'; props.subject='M.Sc. Computer Science Research Project'
    doc.save(DOCX)
    print(DOCX)


if __name__ == '__main__': main()
